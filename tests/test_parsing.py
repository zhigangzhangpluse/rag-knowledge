"""用小文档验证解析和窗口边界，防止正文丢失。"""
import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.parsing import Block, parse, split_blocks
from app.retrieval import rrf


def test_markdown_hierarchy_and_code(tmp_path):
    path = tmp_path / "sample.md"
    path.write_text("# 手册\n## 安装\n请先安装工具。\n```python\nprint('hello')\n```\n## 启动\n执行启动命令。", encoding="utf-8")
    blocks = parse(path)
    assert blocks[0].heading == "手册 > 安装"
    assert "print('hello')" in blocks[1].text
    assert blocks[-1].heading == "手册 > 启动"


def test_docx_preserves_heading_and_table(tmp_path):
    doc = Document()
    doc.add_heading("报销制度", 1)
    doc.add_paragraph("需要提交发票。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "项目", "限额"
    table.cell(1, 0).text, table.cell(1, 1).text = "交通", "100 元"
    path = tmp_path / "policy.docx"
    doc.save(path)
    blocks = parse(path)
    assert blocks[0].heading == "报销制度"
    assert "项目 | 限额" in blocks[1].text and "100 元" in blocks[1].text


def test_scanned_pdf_has_clear_error(tmp_path):
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    path = tmp_path / "scan.pdf"
    writer.write(path)
    with pytest.raises(ValueError, match="OCR"):
        parse(path)


def test_windows_cover_text_and_overlap():
    text = "".join(chr(0x4E00 + i) for i in range(300))
    pieces = split_blocks([Block(text, "章节", "第 1 页")], 80, 20)
    assert all(len(piece["text"]) <= 80 for piece in pieces)
    assert pieces[0]["text"][-20:] == pieces[1]["text"][:20]
    restored = pieces[0]["text"] + "".join(piece["text"][20:] for piece in pieces[1:])
    assert restored == text


def test_rrf_rewards_agreement():
    ranking = rrf([["a", "b"], ["b", "c"]])
    assert ranking[0][0] == "b"
    assert ranking[0][1] == pytest.approx(1 / 62 + 1 / 61)


def test_pdf_text_and_numbered_heading(tmp_path):
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 50 780 Td (1. Leave Policy) Tj 0 -20 Td (1.1 Annual leave) Tj 0 -20 Td (Employees receive five days.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    path = tmp_path / "text.pdf"
    writer.write(path)
    blocks = parse(path)
    assert blocks[0].heading == "1. Leave Policy > 1.1 Annual leave"
    assert "five days" in blocks[0].text
    assert blocks[0].location == "第 1 页"


def test_markdown_table_stays_with_headers(tmp_path):
    path = tmp_path / "table.md"
    path.write_text("# 费用\n\n| 项目 | 上限 |\n| --- | --- |\n| 住宿 | 500 元 |\n", encoding="utf-8")
    blocks = parse(path)
    assert len(blocks) == 1
    assert "项目" in blocks[0].text and "500 元" in blocks[0].text
